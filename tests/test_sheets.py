"""Spreadsheet cell reading and writing.

Skipped when openpyxl is missing, so the rest of the suite still runs on a machine
without it.
"""

import tempfile
import unittest
from pathlib import Path

from phush import sheets
from phush.sheets import SheetError

try:
    import openpyxl

    HAVE_OPENPYXL = True
except ImportError:
    HAVE_OPENPYXL = False


@unittest.skipUnless(HAVE_OPENPYXL, "openpyxl is not installed")
class SheetCases(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.root = Path(self._tmp.name)
        self.book_path = self.root / "table.xlsx"

        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Processors"
        sheet["A1"] = "model"
        sheet["B1"] = "cores"
        sheet["A2"] = "Ampere Altra"
        sheet["B2"] = 80
        sheet["A3"] = "Graviton 3"
        sheet["B3"] = 64
        book.create_sheet("Notes")
        book.save(self.book_path)

    def tearDown(self):
        self._tmp.cleanup()

    def reload(self):
        return openpyxl.load_workbook(self.book_path)

    # ---- reading ----

    def test_describe_lists_sheets_and_size(self):
        out = sheets.describe(self.book_path)
        self.assertIn("Processors", out)
        self.assertIn("Notes", out)
        self.assertIn("3 rows", out)

    def test_read_a_single_cell(self):
        self.assertIn("Ampere Altra", sheets.read_range(self.book_path, "Processors", "A2"))

    def test_read_a_range_with_addresses(self):
        out = sheets.read_range(self.book_path, "Processors", "A1:B2")
        self.assertIn("A1\tmodel", out)
        self.assertIn("B2\t80", out)

    def test_read_lowercase_range_is_accepted(self):
        self.assertIn("A1", sheets.read_range(self.book_path, "Processors", "a1:b1"))

    def test_read_empty_range_says_so(self):
        self.assertIn("empty", sheets.read_range(self.book_path, "Notes", "A1:B2"))

    def test_bad_range_is_refused(self):
        for bad in ["A", "1", "A1:", "Q5; DROP", ""]:
            with self.subTest(ref=bad), self.assertRaises(SheetError):
                sheets.read_range(self.book_path, "Processors", bad)

    def test_unknown_sheet_lists_the_real_ones(self):
        with self.assertRaises(SheetError) as caught:
            sheets.read_range(self.book_path, "Nope", "A1")
        self.assertIn("Processors", str(caught.exception))

    def test_empty_sheet_name_uses_the_first_sheet(self):
        self.assertIn("model", sheets.read_range(self.book_path, "", "A1"))

    def test_missing_file_is_refused(self):
        with self.assertRaises(SheetError):
            sheets.describe(self.root / "nope.xlsx")

    def test_non_spreadsheet_is_refused(self):
        other = self.root / "notes.md"
        other.write_text("hello", encoding="utf-8")
        with self.assertRaises(SheetError):
            sheets.describe(other)

    # ---- searching ----

    def test_find_returns_addresses(self):
        self.assertIn("A3", sheets.find(self.book_path, "Processors", "graviton"))

    def test_find_without_matches_says_so(self):
        self.assertEqual("no matches", sheets.find(self.book_path, "Processors", "zzz"))

    # ---- writing ----

    def test_set_cells_writes_text_and_numbers(self):
        sheets.set_cells(self.book_path, "Processors", {"C1": "note", "C2": "42"})
        sheet = self.reload()["Processors"]
        self.assertEqual("note", sheet["C1"].value)
        self.assertEqual(42, sheet["C2"].value)
        self.assertIsInstance(sheet["C2"].value, int)

    def test_set_cells_keeps_a_formula(self):
        sheets.set_cells(self.book_path, "Processors", {"C2": "=B2*2"})
        self.assertEqual("=B2*2", self.reload()["Processors"]["C2"].value)

    def test_set_cells_leaves_other_cells_alone(self):
        sheets.set_cells(self.book_path, "Processors", {"C2": "x"})
        sheet = self.reload()["Processors"]
        self.assertEqual("Ampere Altra", sheet["A2"].value)
        self.assertEqual(64, sheet["B3"].value)

    def test_set_cells_refuses_a_range_or_junk(self):
        for bad in ["A1:B2", "sheet!A1", "", "A", "../etc"]:
            with self.subTest(address=bad), self.assertRaises(SheetError):
                sheets.set_cells(self.book_path, "Processors", {bad: "x"})

    def test_set_cells_refuses_an_empty_request(self):
        with self.assertRaises(SheetError):
            sheets.set_cells(self.book_path, "Processors", {})

    def test_set_cells_on_unknown_sheet_changes_nothing(self):
        with self.assertRaises(SheetError):
            sheets.set_cells(self.book_path, "Nope", {"A1": "x"})
        self.assertEqual("model", self.reload()["Processors"]["A1"].value)


if __name__ == "__main__":
    unittest.main()
